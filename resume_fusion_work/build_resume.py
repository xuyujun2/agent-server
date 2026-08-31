from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK


OUT = r"C:\Users\28763\Desktop\agent\customer_service_agent\徐玉俊-AI Agent-Python后端开发工程师简历.docx"

NAVY = RGBColor(31, 78, 121)
DARK = RGBColor(34, 34, 34)
GRAY = RGBColor(95, 95, 95)
LIGHT = "D9E2F3"


def set_run_font(run, size=10, bold=False, color=DARK, name="微软雅黑"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=50, start=80, bottom=50, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_bottom_border(paragraph, color="4472C4", size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_section(doc, title):
    p = doc.add_paragraph(style="Resume Heading")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run_font(r, 12.5, True, NAVY)
    add_bottom_border(p)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.03
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, 9.0)
    return p


def add_project(doc, title, tech, bullets, tag="个人项目"):
    p = doc.add_paragraph(style="Project Title")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run_font(r, 10.2, True, DARK)
    r2 = p.add_run(f"  |  {tag}")
    set_run_font(r2, 8.8, False, GRAY)
    for item in bullets:
        add_bullet(doc, item)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(3)
    p2.paragraph_format.keep_together = True
    a = p2.add_run("技术栈：")
    set_run_font(a, 8.8, True, NAVY)
    b = p2.add_run(tech)
    set_run_font(b, 8.8, False, GRAY)


def add_job(doc, company, dates, role, projects, tech_stack):
    p = doc.add_paragraph(style="Job Title")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(company)
    set_run_font(r, 10.5, True, DARK)
    r = p.add_run(f"  |  {role}  |  {dates}")
    set_run_font(r, 9.2, False, GRAY)
    for project in projects:
        title, desc = project[0], project[1]
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Cm(0.2)
        p2.paragraph_format.space_after = Pt(1.5)
        p2.paragraph_format.keep_together = True
        a = p2.add_run(f"{title}：")
        set_run_font(a, 9.0, True)
        b = p2.add_run(desc)
        set_run_font(b, 9.0)
    p3 = doc.add_paragraph()
    p3.paragraph_format.left_indent = Cm(0.2)
    p3.paragraph_format.space_after = Pt(1.2)
    p3.paragraph_format.keep_together = True
    a = p3.add_run("技术栈：")
    set_run_font(a, 8.2, True, NAVY)
    b = p3.add_run(tech_stack)
    set_run_font(b, 8.2, False, GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.0)
sec.bottom_margin = Cm(1.0)
sec.left_margin = Cm(1.55)
sec.right_margin = Cm(1.55)
sec.header_distance = Cm(0.6)
sec.footer_distance = Cm(0.6)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "微软雅黑"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(9.0)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.03

for name in ["Resume Heading", "Project Title", "Job Title"]:
    if name not in styles:
        styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
styles["Resume Heading"].paragraph_format.space_before = Pt(6)
styles["Resume Heading"].paragraph_format.space_after = Pt(4)
styles["Project Title"].paragraph_format.space_before = Pt(3)
styles["Project Title"].paragraph_format.space_after = Pt(1.5)
styles["Job Title"].paragraph_format.space_before = Pt(3)
styles["Job Title"].paragraph_format.space_after = Pt(2)

# Header block
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("徐玉俊")
set_run_font(r, 22, True, NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
r = p.add_run("AI Agent / Python 后端开发工程师（Web 前端背景）")
set_run_font(r, 11.5, True, DARK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
for idx, text in enumerate([
    "15321875083", "aa349593361@sina.com", "北京昌平区龙锦苑", "7年以上 Web 前端经验"
]):
    if idx:
        rr = p.add_run("  |  ")
        set_run_font(rr, 9, False, GRAY)
    rr = p.add_run(text)
    set_run_font(rr, 9.2, False, GRAY)

add_section(doc, "个人概述")
add_bullet(doc, "软件工程本科，拥有7年以上 Web 前端开发经验，熟悉 Vue、React、TypeScript、工程化及复杂业务系统开发；近期系统学习 Python、FastAPI、LangChain 与大模型应用开发。")
add_bullet(doc, "独立完成智能客服、自动批改、代码生成与单元测试 MCP、PDF 报告生成、邮件自动回复等 AI Agent 项目，覆盖需求分析、模型调用、工具设计、接口开发、数据库及 Docker 部署。")
add_bullet(doc, "具备前后端联调和阿里云部署经验，能够将 AI 能力封装为 REST API、SSE 服务或 MCP Tool，并结合真实业务流程完成可运行方案。")

add_section(doc, "专业技能")
skills = [
    ("AI Agent", "LangChain、AgentExecutor、Tool Calling、Prompt Engineering、上下文记忆、DeepSeek API、OpenAI 兼容接口、阿里云百炼视觉模型"),
    ("Python 后端", "Python、FastAPI、Pydantic、RESTful API、SSE、PyMySQL、MySQL、定时任务"),
    ("前端开发", "JavaScript、TypeScript、Vue 2/3、React、Next.js、Pinia、Redux、UniApp、HTML、CSS、ECharts"),
    ("工程与部署", "Git、Vite、Webpack、Docker、Dockerfile、Nginx、Linux、阿里云 ECS、环境变量管理"),
    ("测试与协议", "pytest、MCP、stdio MCP、Streamable HTTP MCP、Gmail API、OAuth 2.0"),
    ("文档与多模态", "PDF 解析、图片 Base64、Pillow、视觉识别、结构化 JSON 处理"),
]
for label, value in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.2)
    a = p.add_run(f"{label}：")
    set_run_font(a, 9.0, True, NAVY)
    b = p.add_run(value)
    set_run_font(b, 9.0)

add_section(doc, "AI Agent 项目经历")
add_project(doc, "智能客服 Agent 平台", "Python、FastAPI、LangChain、DeepSeek、MySQL、PyMySQL、Docker、Nginx、SSE", [
    "基于 FastAPI、LangChain 与 DeepSeek 构建智能客服系统，将订单查询、退货申请等业务封装为 Tools，由 Agent 根据用户问题自主选择并执行。",
    "设计订单、退货申请等 MySQL 表结构，通过 PyMySQL 完成查询与更新；实现会话历史和上下文记忆，支持多轮问答。",
    "实现 SSE 流式响应和支付状态回调；通过 Docker 部署 API 与 MySQL，并完成阿里云 ECS、前端及 Nginx 联调。",
])
add_project(doc, "AI 自动批改作业系统", "Python、FastAPI、LangChain、DeepSeek、阿里云百炼视觉模型、Pillow、JSON、Base64", [
    "设计多文件上传接口，支持标准答案与学生答题卷的多图上传；使用视觉模型识别题号、题目、答案及原图坐标。",
    "将识别结果和评分规则提交给批改 Agent，分别处理数学与作文场景，输出逐题得分、正误判断和批改说明。",
    "针对多页试卷进行识别文本合并及跨页答案关联适配；使用程序汇总逐题分数，并用 Pillow 将勾叉、得分和说明绘制回原图。",
])
add_project(doc, "代码生成与单元测试 MCP 服务", "Python、MCP、FastMCP、LangChain、DeepSeek、pytest、Docker、FastAPI", [
    "基于大模型生成 Python 业务代码，并根据执行错误进行多轮修正；自动生成并运行正常、边界和异常场景 pytest 测试。",
    "实现 stdio 与 Streamable HTTP 两种 MCP Server，注册代码生成和测试工具，并完成与 Codex 的连接验证。",
    "使用独立 Docker 沙箱执行生成代码，限制网络、内存、CPU、进程数、执行时间和文件权限；通过环境变量管理 API Key。",
])
add_project(doc, "PDF 智能解析与报告生成", "Python、FastAPI、pdfplumber、DeepSeek、LangChain", [
    "实现 PDF 上传、文本提取、分段与结构化分析，调用大模型生成摘要、关键结论和分析报告。",
    "将 PDF 解析及报告生成逻辑封装为独立 Service，通过 FastAPI 对外提供能力。",
])
add_project(doc, "智能邮件自动回复 Agent", "Python、Gmail API、OAuth 2.0、DeepSeek、定时任务、Docker", [
    "接入 Gmail API 和 OAuth 2.0，定时扫描未读邮件，解析主题、发件人及纯文本或 HTML 正文，并处理邮件头中文编码。",
    "按白名单和主题关键词匹配自动回复规则，由 DeepSeek 生成回复并通过 Gmail API 发送；处理完成后移除未读标签，避免重复执行。",
    "通过 Docker 文件挂载提供 OAuth Token，避免将授权文件打入镜像。",
])

add_section(doc, "工作经历")
add_job(doc, "北京新美互通科技有限公司", "2025.05 - 2026.03", "Web 前端工程师", [
    ("虚拟人 H5（Chatlyverse）", "负责 Vue 3 项目搭建及全页面开发，实现 SSE AI 聊天、Firebase 登录与埋点、Google 广告接入、rem 适配、性能与 SEO 优化。"),
    ("优惠券 H5（Couponhubx）", "负责项目搭建、品牌优惠券流程、广告接入、Firebase 埋点、移动端适配及加载性能优化。"),
], "Vue 3、SSE、Pinia、Firebase、Fetch、Storage、Google AdSense、Google Ad Manager")
add_job(doc, "北京点众科技股份有限公司", "2022.03 - 2025.04", "Web 前端工程师", [
    ("小说大全", "负责阅读页排版与分页、翻页动画、付费章节处理；参与短剧滑动播放、缓存优化、会员收银台、支付及书城模块开发。"),
    ("嵌入式 H5 活动平台", "负责活动页面与性能优化，完成大转盘、抽奖、充值活动、优惠券、分享落地页等功能。"),
    ("河马免费小说", "负责日常迭代及广告业务开发，包括广告结构调整、交互方案与摇一摇等能力。"),
], "Vue、Vue Router、Axios、Storage、Vant、Vite、Video、Swiper、微信支付、支付宝、神策埋点")
add_job(doc, "中文未来教育科技有限公司", "2019.06 - 2022.02", "Web 前端工程师", [
    ("会员卡销售系统", "负责首页 ECharts 数据统计、会员卡配置与审核、权限分配、代理商及分润管理。"),
    ("校园招聘系统", "使用 React、Next.js、Redux、TypeScript 开发简历、应聘信息、笔试试题、试卷及应聘结果管理。"),
    ("快解阅读小程序", "使用 UniApp 开发商城与个人中心，涵盖学员信息、激活码、答题记录、学习币及支付流程。"),
], "Vue、Vuex、Pinia、Vue Router、Axios、ECharts、Element Plus、React Hooks、Redux、Next.js、Ant Design、TypeScript、UniApp、Vant Weapp")
add_job(doc, "征和控股有限公司", "2018.06 - 2019.06", "Web 前端工程师", [
    ("多交易所项目", "负责登录注册、首页及公共组件、帮助中心、账户设置、风险评测等 Vue 业务模块。"),
    ("借款人项目", "负责内部数据平台的手续费、服务费、充值记录和提现记录等模块。"),
], "Vue CLI、Vue、Vuex、Vue Router、Axios、LocalStorage、Element UI、按需加载")

add_section(doc, "教育背景")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("南京财经大学")
set_run_font(r, 10.5, True)
r = p.add_run("  |  软件工程  |  本科  |  2014.09 - 2018.06")
set_run_font(r, 9.5, False, GRAY)

add_section(doc, "个人评价")
add_bullet(doc, "持续学习并主动将新技术落地为可运行项目，具备较强的逻辑分析、问题定位和独立开发能力。")
add_bullet(doc, "拥有多年团队协作与业务交付经验，沟通清晰、责任心强，重视代码质量、性能优化和用户体验。")

# Footer
for section in doc.sections:
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run("徐玉俊｜AI Agent / Python 后端开发工程师")
    set_run_font(rr, 8, False, GRAY)

doc.core_properties.title = "徐玉俊 - AI Agent / Python 后端开发工程师简历"
doc.core_properties.subject = "AI Agent、Python 后端、Web 前端"
doc.core_properties.author = "徐玉俊"
doc.save(OUT)
print(OUT)
